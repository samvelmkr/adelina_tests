# pip install -r requirements.txt

import docx
from docx.shared import RGBColor
import os, sys, re
import openpyxl

q_rows = []


def process_tests(lines) -> None:
  test = []

  for line in lines:
    if not line:
      if test:
        answ = "\n".join(test[1:])
        q_rows.append([test[0], answ])
      test = []
      continue

    test.append(line)


def main():
  if len(sys.argv) < 2:
    print("Usage: python myscript.py <docx filename>")
    sys.exit()

  filename = sys.argv[1]
  assert os.path.exists(filename)

  lines = []
  doc = docx.Document(filename)
  for para in doc.paragraphs:
    line = para.text.strip()
    lines.append(line)

  process_tests(lines)

  workbook = openpyxl.Workbook()

  sheet = workbook.active
  sheet.title = "Sheet"

  for q_row in q_rows:
    # print(q_row)
    sheet.append(q_row)

  filename += ".xlsx"
  workbook.save(filename)


if __name__ == '__main__':
    main()
